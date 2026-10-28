"""Long-format CSV / JSON / XLSX lab-result template parse and download."""
from __future__ import annotations

import csv
import json
import os
import re
from datetime import date, datetime
from io import BytesIO, StringIO
from typing import Any, Dict, List, Optional, Sequence, Tuple

from fastapi import APIRouter, Depends, Request
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from .database import get_db
from .lab_matching import is_plain_number, match_analyte
from .panel_catalog import resolve_household_id

router = APIRouter()

TEMPLATE_HEADERS = [
    "collection_date",
    "panel_name",
    "original_name",
    "value",
    "unit",
    "ref_low",
    "ref_high",
    "ref_text",
    "flag",
    "performing_lab",
    "specimen",
]
REQUIRED_HEADERS = {"collection_date", "panel_name", "original_name", "value"}
FORMAT_NAME = "phr-lab-results"
TABULAR_SUFFIXES = {".csv", ".xls", ".xlsx", ".json"}
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLS_MIME = "application/vnd.ms-excel"

STRUCTURED_MISS_MESSAGE = (
    "The file did not match the lab results template and no result rows were detected. "
    "Use the long template (one result per row) for several panels and dates in one file."
)
XLS_STORE_MESSAGE = (
    "Excel 97-2003 .xls is stored; save as .xlsx for structured import."
)
WIDE_TEMPLATE_MESSAGE = (
    "This spreadsheet looks like dates as columns. Use the long lab results template "
    "(one result per row)."
)

DATE_HEADER_RE = re.compile(
    r"^\s*(\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{2,4})\s*$"
)

EXAMPLE_RESULTS: List[Dict[str, Any]] = [
    {
        "collection_date": "2026-09-21",
        "panel_name": "CMP",
        "original_name": "Sodium",
        "value": "140",
        "unit": "mmol/L",
        "ref_low": 135,
        "ref_high": 145,
        "ref_text": "135-145",
        "flag": None,
        "performing_lab": "Quest",
        "specimen": "serum",
    },
    {
        "collection_date": "2026-09-21",
        "panel_name": "CBC",
        "original_name": "WBC",
        "value": "6.2",
        "unit": "10*3/uL",
        "ref_low": None,
        "ref_high": None,
        "ref_text": None,
        "flag": None,
        "performing_lab": "Quest",
        "specimen": None,
    },
    {
        "collection_date": "2026-05-01",
        "panel_name": "BMP",
        "original_name": "Sodium",
        "value": "138",
        "unit": "mmol/L",
        "ref_low": 135,
        "ref_high": 145,
        "ref_text": "135-145",
        "flag": None,
        "performing_lab": None,
        "specimen": None,
    },
]


def _suffix(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def is_tabular_upload(filename: str, mime: str) -> bool:
    mime = (mime or "").lower()
    suf = _suffix(filename)
    if suf in TABULAR_SUFFIXES:
        return True
    return mime in (
        "text/csv",
        "application/csv",
        "application/json",
        "text/json",
        XLSX_MIME,
        XLS_MIME,
        "application/vnd.ms-excel",
    )


def is_legacy_xls(filename: str, mime: str) -> bool:
    suf = _suffix(filename)
    if suf == ".xls":
        return True
    if suf:
        return False
    return (mime or "").lower() == XLS_MIME


def _as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_collection_date(value: Any) -> Optional[datetime]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Plain numbers are not collection dates (Excel serials must already be datetime).
        return None
    text = _as_text(value)
    if not text:
        return None
    text = re.sub(r"[ T]\d{1,2}:\d{2}.*$", "", text)
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _parse_float(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _as_text(value)
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _norm_headers(headers: Sequence[Any]) -> Dict[str, int]:
    out: Dict[str, int] = {}
    for i, raw in enumerate(headers):
        key = _as_text(raw).strip().lower()
        if key and key not in out:
            out[key] = i
    return out


def _headers_match_template(header_map: Dict[str, int]) -> bool:
    return REQUIRED_HEADERS.issubset(header_map.keys())


def _headers_look_wide(headers: Sequence[Any]) -> bool:
    names = [_as_text(h) for h in headers if _as_text(h)]
    if not names:
        return False
    lower = {n.lower() for n in names}
    if REQUIRED_HEADERS.issubset(lower):
        return False
    date_like = sum(1 for n in names if DATE_HEADER_RE.match(n))
    return date_like >= 2


def _decode_text(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="replace")


def _row_from_mapping(mapping: Dict[str, Any], analytes: list, quote_fallback: str) -> Optional[Dict[str, Any]]:
    keys = {k.strip().lower(): k for k in mapping.keys() if isinstance(k, str)}

    def get(name: str) -> Any:
        src = keys.get(name)
        if src is None:
            return None
        return mapping.get(src)

    original = _as_text(get("original_name"))
    value = get("value")
    value_text = _as_text(value)
    panel = _as_text(get("panel_name"))
    if not original or value_text == "" or not panel:
        return None
    unit = _as_text(get("unit")) or None
    coll = parse_collection_date(get("collection_date"))
    matched = match_analyte(original, analytes) if analytes is not None else None
    flag = _as_text(get("flag")) or None
    performing = _as_text(get("performing_lab")) or None
    specimen = _as_text(get("specimen")) or None
    date_label = coll.strftime("%Y-%m-%d") if coll else _as_text(get("collection_date"))
    quote = f"{panel} {original} {value_text} {unit or ''} on {date_label}".strip()
    if not quote.strip():
        quote = quote_fallback
    return {
        "original_name": original,
        "value_text": value_text,
        "unit_original": unit,
        "ref_low": _parse_float(get("ref_low")),
        "ref_high": _parse_float(get("ref_high")),
        "ref_text": _as_text(get("ref_text")) or None,
        "printed_flag": flag,
        "collection_date": coll,
        "panel_name": panel,
        "performing_lab": performing,
        "specimen": specimen,
        "matched": bool(matched),
        "analyte_id": matched.id if matched else None,
        "plain_number": is_plain_number(value_text),
        "quote": quote,
    }


def _parse_json_obj(obj: Any) -> Optional[List[Dict[str, Any]]]:
    if not isinstance(obj, dict):
        return None
    fmt = obj.get("format")
    results = obj.get("results")
    if fmt == FORMAT_NAME and isinstance(results, list):
        return [r for r in results if isinstance(r, dict)]
    if "version" in obj and isinstance(results, list):
        if not results:
            return []
        if all(isinstance(r, dict) and "original_name" in r and "value" in r for r in results):
            return results
    return None


def looks_like_structured_json(data: bytes) -> bool:
    try:
        obj = json.loads(_decode_text(data))
    except (ValueError, UnicodeDecodeError):
        return False
    return _parse_json_obj(obj) is not None


def looks_like_structured_csv(data: bytes) -> bool:
    try:
        text = _decode_text(data)
        reader = csv.reader(StringIO(text))
        header = next(reader, None)
    except Exception:
        return False
    if not header:
        return False
    return _headers_match_template(_norm_headers(header))


def _xlsx_header_row(data: bytes) -> Optional[List[Any]]:
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return None
    try:
        sheet = wb["results"] if "results" in wb.sheetnames else wb[wb.sheetnames[0]]
        for row in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
            return list(row or [])
        return []
    finally:
        wb.close()


def looks_like_structured_xlsx(data: bytes) -> bool:
    header = _xlsx_header_row(data)
    if header is None:
        return False
    return _headers_match_template(_norm_headers(header))


def looks_like_wide_lab(filename: str, mime: str, data: bytes) -> bool:
    suf = _suffix(filename)
    if suf == ".csv" or (mime or "").lower() in ("text/csv", "application/csv"):
        try:
            header = next(csv.reader(StringIO(_decode_text(data))), None)
        except Exception:
            return False
        return bool(header) and _headers_look_wide(header)
    if suf == ".xlsx" or (mime or "").lower() == XLSX_MIME:
        header = _xlsx_header_row(data)
        return bool(header) and _headers_look_wide(header)
    return False


def looks_like_structured_lab(filename: str, mime: str, data: bytes) -> bool:
    suf = _suffix(filename)
    mime_l = (mime or "").lower()
    if suf == ".json" or mime_l in ("application/json", "text/json"):
        return looks_like_structured_json(data)
    if suf == ".csv" or mime_l in ("text/csv", "application/csv"):
        return looks_like_structured_csv(data)
    if suf == ".xlsx" or mime_l == XLSX_MIME:
        return looks_like_structured_xlsx(data)
    return False


def parse_structured_json(data: bytes, analytes: list) -> List[Dict[str, Any]]:
    obj = json.loads(_decode_text(data))
    raw_rows = _parse_json_obj(obj) or []
    out = []
    for i, mapping in enumerate(raw_rows):
        row = _row_from_mapping(mapping, analytes, quote_fallback=f"json results[{i}]")
        if row:
            out.append(row)
    return out


def parse_structured_csv(data: bytes, analytes: list) -> List[Dict[str, Any]]:
    text = _decode_text(data)
    reader = csv.DictReader(StringIO(text))
    out = []
    for i, mapping in enumerate(reader, start=2):
        row = _row_from_mapping(mapping, analytes, quote_fallback=f"csv line {i}")
        if row:
            out.append(row)
    return out


def parse_structured_xlsx(data: bytes, analytes: list) -> List[Dict[str, Any]]:
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = wb["results"] if "results" in wb.sheetnames else wb[wb.sheetnames[0]]
        rows_iter = sheet.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            return []
        header_map = _norm_headers(header)
        out = []
        for i, values in enumerate(rows_iter, start=2):
            mapping: Dict[str, Any] = {}
            for name, idx in header_map.items():
                mapping[name] = values[idx] if idx < len(values) else None
            row = _row_from_mapping(mapping, analytes, quote_fallback=f"xlsx row {i}")
            if row:
                out.append(row)
        return out
    finally:
        wb.close()


def parse_structured_lab(data: bytes, filename: str, analytes: list) -> List[Dict[str, Any]]:
    suf = _suffix(filename)
    if suf == ".json":
        return parse_structured_json(data, analytes)
    if suf == ".csv":
        return parse_structured_csv(data, analytes)
    if suf == ".xlsx":
        return parse_structured_xlsx(data, analytes)
    if looks_like_structured_json(data):
        return parse_structured_json(data, analytes)
    if looks_like_structured_csv(data):
        return parse_structured_csv(data, analytes)
    if looks_like_structured_xlsx(data):
        return parse_structured_xlsx(data, analytes)
    return []


def extract_xlsx_text(data: bytes) -> Tuple[str, int, bool]:
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return "", 0, False
    try:
        names = wb.sheetnames
        sheet = wb["results"] if "results" in names else wb[names[0]]
        lines = []
        for row in sheet.iter_rows(values_only=True):
            cells = []
            for cell in row:
                if cell is None:
                    cells.append("")
                elif isinstance(cell, datetime):
                    cells.append(cell.strftime("%Y-%m-%d"))
                else:
                    cells.append(str(cell))
            lines.append("\t".join(cells).rstrip())
        text = "\n".join(lines).strip()
        return text, 1, False
    except Exception:
        return "", 0, False
    finally:
        wb.close()


def build_template_csv() -> str:
    buf = StringIO()
    writer = csv.DictWriter(buf, fieldnames=TEMPLATE_HEADERS, lineterminator="\n")
    writer.writeheader()
    for row in EXAMPLE_RESULTS:
        writer.writerow({k: ("" if row.get(k) is None else row.get(k)) for k in TEMPLATE_HEADERS})
    return buf.getvalue()


def build_template_json() -> str:
    return json.dumps(
        {"version": 1, "format": FORMAT_NAME, "results": EXAMPLE_RESULTS},
        indent=2,
    )


def build_template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "results"
    ws.append(TEMPLATE_HEADERS)
    for row in EXAMPLE_RESULTS:
        ws.append([row.get(k) for k in TEMPLATE_HEADERS])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/lab-import/template.csv")
def download_template_csv(request: Request, db: Session = Depends(get_db)):
    resolve_household_id(request, db)
    return Response(
        content=build_template_csv(),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=phr-lab-results-template.csv"},
    )


@router.get("/lab-import/template.json")
def download_template_json(request: Request, db: Session = Depends(get_db)):
    resolve_household_id(request, db)
    return Response(
        content=build_template_json(),
        media_type="application/json",
        headers={"Content-Disposition": "attachment; filename=phr-lab-results-template.json"},
    )


@router.get("/lab-import/template.xlsx")
def download_template_xlsx(request: Request, db: Session = Depends(get_db)):
    resolve_household_id(request, db)
    return Response(
        content=build_template_xlsx(),
        media_type=XLSX_MIME,
        headers={"Content-Disposition": "attachment; filename=phr-lab-results-template.xlsx"},
    )

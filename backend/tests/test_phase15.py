"""Phase 15: structured CSV / XLS / XLSX / JSON lab import."""
import os

os.environ["DATABASE_URL"] = "sqlite:////tmp/test_phase15.db"
os.environ["STORAGE_ROOT"] = "/tmp/phr_docs_test15"
os.environ.setdefault("AUTH_DISABLED", "1")

import json
import shutil
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import get_db
from app.extraction import MAX_UPLOAD_BYTES
from app.main import app
from app.models import Base, CanonicalAnalyte, Document, LabReport

FIXTURES = Path(__file__).parent / "fixtures"

test_engine = create_engine(
    "sqlite:////tmp/test_phase15.db",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


client = TestClient(app)


@pytest.fixture(autouse=True)
def clean_db():
    app.dependency_overrides[get_db] = override_get_db
    Base.metadata.drop_all(bind=test_engine)
    Base.metadata.create_all(bind=test_engine)
    db = TestingSessionLocal()
    from app.seed import load_analytes

    load_analytes(db)
    db.close()
    shutil.rmtree("/tmp/phr_docs_test15", ignore_errors=True)
    os.makedirs("/tmp/phr_docs_test15", exist_ok=True)
    yield


def seed_count() -> int:
    db = TestingSessionLocal()
    try:
        return db.query(CanonicalAnalyte).count()
    finally:
        db.close()


def new_profile(name="Test"):
    return client.post("/profiles", json={"legal_name": name}).json()["id"]


def process_upload(pid: int, filename: str, data: bytes, mime: str):
    r = client.post(
        f"/profiles/{pid}/documents",
        files={"file": (filename, data, mime)},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    pr = client.post(f"/jobs/{body['job_id']}/process")
    assert pr.status_code == 200, pr.text
    docs = client.get(f"/profiles/{pid}/documents").json()
    doc = next(d for d in docs if d["id"] == body["document_id"])
    job = {"status": pr.json()["status"], "error_message": doc.get("error")}
    return body, job


def load_fixture(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


def make_xlsx_from_csv_bytes(csv_bytes: bytes) -> bytes:
    import csv
    from io import StringIO

    text = csv_bytes.decode("utf-8-sig")
    reader = csv.DictReader(StringIO(text))
    headers = list(reader.fieldnames or [])
    wb = Workbook()
    ws = wb.active
    ws.title = "results"
    ws.append(headers)
    for row in reader:
        ws.append([row.get(h, "") for h in headers])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def assert_multi_panel_multi_date(pid: int):
    labs = client.get(f"/profiles/{pid}/lab-reports").json()
    assert len(labs) == 3
    by_panel = {lab["panel_name"]: lab for lab in labs}
    assert set(by_panel) == {"CMP", "CBC", "BMP"}
    cmp_names = [c["original_name"] for c in by_panel["CMP"]["components"]]
    cbc_names = [c["original_name"] for c in by_panel["CBC"]["components"]]
    bmp_names = [c["original_name"] for c in by_panel["BMP"]["components"]]
    assert cmp_names == ["Sodium"]
    assert cbc_names == ["WBC"]
    assert bmp_names == ["Sodium"]
    encounters = client.get(f"/profiles/{pid}/encounters").json()
    lab_enc = [e for e in encounters if e["type"] == "lab"]
    assert len(lab_enc) == 2
    starts = sorted(e["start"][:10] for e in lab_enc)
    assert starts == ["2026-05-01", "2026-09-21"]
    assert (by_panel["CMP"]["collection_time"] or "")[:10] == "2026-09-21"
    assert (by_panel["CBC"]["collection_time"] or "")[:10] == "2026-09-21"
    assert (by_panel["BMP"]["collection_time"] or "")[:10] == "2026-05-01"
    na = client.get(f"/profiles/{pid}/trends/Sodium").json()
    assert len(na["points"]) == 2
    na_dates = sorted((p["collected_at"] or "")[:10] for p in na["points"])
    assert na_dates == ["2026-05-01", "2026-09-21"]
    wbc = client.get(f"/profiles/{pid}/trends/Leukocytes").json()
    assert len(wbc["points"]) == 1
    assert (wbc["points"][0]["collected_at"] or "")[:10] == "2026-09-21"


def test_json_template_files_multi_panel_multi_date():
    before = seed_count()
    pid = new_profile()
    data = load_fixture("lab-import-sample.json")
    _, job = process_upload(pid, "labs.json", data, "application/json")
    assert job["status"] == "filed"
    assert_multi_panel_multi_date(pid)
    assert seed_count() == before


def test_csv_template_files_multi_panel_multi_date():
    pid = new_profile()
    data = load_fixture("lab-import-sample.csv")
    process_upload(pid, "labs.csv", data, "text/csv")
    assert_multi_panel_multi_date(pid)


def test_xlsx_template_files_multi_panel_multi_date():
    pid = new_profile()
    data = make_xlsx_from_csv_bytes(load_fixture("lab-import-sample.csv"))
    process_upload(
        pid,
        "labs.xlsx",
        data,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert_multi_panel_multi_date(pid)


def test_unknown_name_ferritin_goes_to_inbox():
    pid = new_profile()
    body = {
        "version": 1,
        "format": "phr-lab-results",
        "results": [
            {
                "collection_date": "2026-09-21",
                "panel_name": "Iron",
                "original_name": "Ferritin",
                "value": "45",
                "unit": "ng/mL",
            }
        ],
    }
    _, job = process_upload(
        pid, "ferritin.json", json.dumps(body).encode("utf-8"), "application/json"
    )
    assert job["status"] == "needs_review"
    labs = client.get(f"/profiles/{pid}/lab-reports").json()
    assert all("Ferritin" not in c["original_name"] for lab in labs for c in lab["components"])
    props = client.get(f"/profiles/{pid}/proposals").json()
    assert len(props) >= 1
    blob = " ".join((pr.get("quote") or "") + str(pr.get("payload")) for pr in props)
    assert "Ferritin" in blob
    assert "45" in blob


def test_nontemplate_json_is_stored_without_pdf_empty_copy():
    pid = new_profile()
    data = load_fixture("lab-import-nontemplate.json")
    uploaded, job = process_upload(pid, "notes.json", data, "application/json")
    assert job["status"] == "needs_review"
    err = job.get("error_message") or ""
    assert "Sodium 140 mmol/L" not in err
    assert "template" in err.lower()
    text = client.get(f"/documents/{uploaded['document_id']}/text").json()
    extracted = text.get("extracted_text") or ""
    assert extracted.strip()
    assert "foo" in extracted
    docs = client.get(f"/profiles/{pid}/documents").json()
    assert len(docs) == 1
    assert docs[0]["filename"] == "notes.json"


def test_upload_over_50mb_is_413():
    pid = new_profile()
    too_big = b"x" * (MAX_UPLOAD_BYTES + 1)
    r = client.post(
        f"/profiles/{pid}/documents",
        files={"file": ("huge.csv", too_big, "text/csv")},
    )
    assert r.status_code == 413


def test_profile_isolation():
    p1 = new_profile("One")
    p2 = new_profile("Two")
    data = load_fixture("lab-import-sample.json")
    process_upload(p1, "labs.json", data, "application/json")
    assert client.get(f"/profiles/{p2}/documents").json() == []
    assert client.get(f"/profiles/{p2}/lab-reports").json() == []
    assert len(client.get(f"/profiles/{p1}/lab-reports").json()) == 3


def test_octet_stream_json_uses_suffix_mime():
    pid = new_profile()
    data = load_fixture("lab-import-sample.json")
    uploaded, _ = process_upload(pid, "labs.json", data, "application/octet-stream")
    db = TestingSessionLocal()
    try:
        doc = db.query(Document).filter(Document.id == uploaded["document_id"]).first()
        assert "json" in (doc.mime_type or "").lower()
    finally:
        db.close()


def test_legacy_xls_is_stored_with_save_as_xlsx_message():
    pid = new_profile()
    uploaded, job = process_upload(pid, "old.xls", b"not-a-real-xls", "application/vnd.ms-excel")
    assert job["status"] == "needs_review"
    err = (job.get("error_message") or "").lower()
    assert ".xls" in err
    assert "xlsx" in err
    docs = client.get(f"/profiles/{pid}/documents").json()
    assert docs[0]["id"] == uploaded["document_id"]


def test_templates_downloadable():
    csv_r = client.get("/lab-import/template.csv")
    assert csv_r.status_code == 200
    assert "collection_date" in csv_r.text
    assert "panel_name" in csv_r.text
    assert "attachment" in csv_r.headers.get("content-disposition", "").lower()
    json_r = client.get("/lab-import/template.json")
    assert json_r.status_code == 200
    body = json_r.json()
    assert body["format"] == "phr-lab-results"
    assert body["version"] == 1
    assert isinstance(body["results"], list) and len(body["results"]) >= 2
    xlsx_r = client.get("/lab-import/template.xlsx")
    assert xlsx_r.status_code == 200
    wb = load_workbook(BytesIO(xlsx_r.content))
    assert "results" in wb.sheetnames
    headers = [c.value for c in next(wb["results"].iter_rows(min_row=1, max_row=1))]
    assert "collection_date" in headers
    assert "original_name" in headers
    assert "value" in headers
